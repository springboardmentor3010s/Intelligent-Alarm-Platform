"""Build Milestone 3 and Milestone 4 DOCX files mirroring the exact format of
Milestone - 1 and Milestone - 2 (same headings, tables, bullets, and the
5-minute trainer script structure). Uses python-docx with matching styles.
"""
import os

import docx
import docx.shared
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import RGBColor

BLUE = RGBColor(0x1F, 0x4E, 0x79)

OUT_DIR = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "..", "Completed Milestones",
)
OUT_DIR = os.path.abspath(OUT_DIR)


def style_defaults(doc):
    n = doc.styles['Normal']
    n.font.name = 'Calibri'
    n.font.size = docx.shared.Pt(10.5)
    h1 = doc.styles['Heading 1']
    h1.font.color.rgb = BLUE
    h1.font.bold = True
    h1.font.size = docx.shared.Pt(14)
    h2 = doc.styles['Heading 2']
    h2.font.color.rgb = BLUE
    h2.font.bold = True
    h2.font.size = docx.shared.Pt(12)
    # ensure Calibri maps to its ascii/hAnsi fonts
    for s in (n, h1, h2):
        rpr = s.element.get_or_add_rPr()
        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is None:
            rfonts = rpr.makeelement(qn('w:rFonts'), {})
            rpr.append(rfonts)
        rfonts.set(qn('w:ascii'), 'Calibri')
        rfonts.set(qn('w:hAnsi'), 'Calibri')


def set_page(doc):
    sec = doc.sections[0]
    sec.page_width = docx.shared.Emu(7772400)
    sec.page_height = docx.shared.Emu(10058400)
    sec.left_margin = docx.shared.Emu(731520)
    sec.right_margin = docx.shared.Emu(731520)
    sec.top_margin = docx.shared.Emu(640080)
    sec.bottom_margin = docx.shared.Emu(640080)


def para(doc, text, style='Normal'):
    p = doc.add_paragraph(text, style=style)
    return p


def bullet(doc, text):
    return para(doc, text, 'List Bullet')


def heading1(doc, text):
    return doc.add_heading(text, level=1)


def heading2(doc, text):
    return doc.add_heading(text, level=2)


def add_table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    for c, h in enumerate(headers):
        t.rows[0].cells[c].text = h
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            t.rows[r + 1].cells[c].text = val
    doc.add_paragraph()
    return t


def build(milestone, subtitle, project_extra, tech_intro, tech_rows, lib_rows,
          milestone_title, required_work, outcomes, completed_intro, completed_rows,
          proof_smoke, proof_intro, proof_bullets, proof_extra_heading, proof_extra_bullets,
          regression_heading, regression_text, flow_line, flow_steps, flow_point,
          issues, script):
    doc = docx.Document()
    style_defaults(doc)
    set_page(doc)

    para(doc, f"Milestone {milestone} - Presentation Notes")
    para(doc, subtitle)

    heading1(doc, "1. PROJECT - WHAT AM I BUILDING?")
    para(doc, "I am building an Intelligent Cognitive Alarm Platform. The concept is simple: a normal alarm rings and can be dismissed half-asleep. This platform requires the user to solve a short cognitive challenge - like a quick math or logic puzzle - before the alarm can be turned off.")
    for line in project_extra:
        para(doc, line)

    heading1(doc, "2. TECHNOLOGY STACK")
    add_table(doc, ["Layer", "Technology", "What I Use It For"], tech_rows)
    para(doc, tech_intro)
    add_table(doc, ["Library", "What I Use It For"], lib_rows)

    heading1(doc, f"3. WHAT WAS MILESTONE {milestone}?")
    para(doc, milestone_title)
    para(doc, "Required work:")
    for b in required_work:
        bullet(doc, b)
    para(doc, outcomes)

    heading1(doc, f"4. WHAT I ACTUALLY COMPLETED")
    para(doc, completed_intro)
    add_table(doc, ["Area", "What I Completed"], completed_rows)

    heading1(doc, f"5. PROOF THAT MILESTONE {milestone} IS COMPLETE")
    heading2(doc, proof_smoke)
    para(doc, proof_intro)
    for b in proof_bullets:
        bullet(doc, b)
    if proof_extra_heading:
        heading2(doc, proof_extra_heading)
    for b in proof_extra_bullets:
        bullet(doc, b)
    heading2(doc, regression_heading)
    para(doc, regression_text)

    heading1(doc, "6. HOW THE SYSTEM WORKS RIGHT NOW")
    para(doc, flow_line)
    for i, s in enumerate(flow_steps, 1):
        para(doc, f"{i}. {s}")
    para(doc, flow_point)

    heading1(doc, "7. TECHNICAL ISSUES I FIXED")
    for b in issues:
        bullet(doc, b)

    heading1(doc, "8. MY 5-MINUTE TRAINER SCRIPT")
    for sub, paras in script:
        heading2(doc, sub)
        for t in paras:
            para(doc, t)

    out = os.path.join(OUT_DIR, f"Milestone - {milestone}.docx")
    doc.save(out)
    print("Saved:", out)


# ---------------------------------------------------------------------------
# MILESTONE 3 : Adaptive Intelligence & Recommendations (Weeks 5-6)
# ---------------------------------------------------------------------------
build(
    3,
    "Intelligent Cognitive Alarm Platform  |  Infosys Virtual Internship  |  Milestone 3 (Weeks 5-6)",
    [
        "Milestone 1 built the foundation: authentication, roles, user profile and settings, timezone-aware alarms, and basic habit scores. Milestone 2 made the alarm require a real cognitive challenge, verified by the server, before it can be dismissed.",
        "In Milestone 3, the platform becomes smarter. The goal was to take the verified wake-up data we now collect and turn it into adaptive intelligence - the system learns each user's wake-up behaviour, adapts the difficulty of the cognitive challenges, scores habits in a more meaningful way, and gives personalised recommendations for sleep and wake-up times.",
    ],
    "Important libraries and frameworks actually used at this stage:",
    [
        ("Main Language", "Python", "Core backend / application logic"),
        ("Backend", "FastAPI", "API and server-side intelligence logic"),
        ("Web Frontend", "React.js", "Web dashboard wired to live data"),
        ("Mobile", "React Native (Expo)", "Mobile alarm/challenge interface"),
        ("Database", "PostgreSQL", "Persistent application data"),
        ("Authentication", "JWT / RBAC", "Secure authentication and authorization"),
        ("Machine Learning", "PyTorch", "Adaptive difficulty model"),
    ],
    [
        ("PyTorch", "Training and running the adaptive difficulty model"),
        ("scikit-learn / XGBoost", "Present in the environment for ML workflows"),
        ("SQLAlchemy", "ORM for models and analytics queries"),
        ("Pydantic", "Request/response validation"),
        ("PyJWT", "Signing and verifying JWT tokens"),
        ("bcrypt", "Password hashing for user registration"),
        ("SQLite", "Local development database (PostgreSQL in Docker)"),
    ],
    "Milestone 3 - Adaptive Intelligence & Recommendations (Weeks 5-6).",
    [
        "Implement adaptive difficulty engine",
        "Build behavioral analytics workflows",
        "Develop habit scoring models",
        "Generate recommendation engine outputs",
        "Create habit tracking dashboards",
    ],
    "Expected outcomes: adaptive difficulty engine operational, habit scoring system functional, and the recommendation workflow completed.",
    "The table below lists what was actually implemented and verified for Milestone 3. The important idea is that every value shown to the user now comes from real data and real behaviour, not placeholder values.",
    [
        ("Analytics Service", "Refactored the analytics layer so sleep duration, wake-up consistency, and trends are computed from real alarm and challenge data instead of placeholders."),
        ("Adaptive Difficulty Model", "Trained a PyTorch DifficultyPredictor (backend/ml_models/difficulty_model.pt) against a synthetic rule-based dataset; the model is loaded and used to pick each challenge's difficulty."),
        ("ML Status Check", "Added GET /api/v1/analytics/ml-status - the endpoint reports model_loaded=true and confirms inference works end to end."),
        ("Rule-Based Fallback", "A deterministic fallback keeps difficulty selection working even when the heavy PyTorch dependency is not available, so a lean runtime still functions."),
        ("Recommendation Engine", "Implemented GET /api/v1/analytics/recommendations returning four categories - sleep, wake-up, habit, and challenge - sorted by priority."),
        ("Sleep & Wake-Up Coaching", "Recommendations suggest 15-minute adjustments to bedtime and wake-up time based on the user's sleep window and preferences."),
        ("Habit Improvement Guidance", "Behavioural thresholds (for example a 20% completion drop) trigger coaching messages instead of static tips."),
        ("Challenge Recommendations", "Recommends unused challenge types and the next difficulty level based on the user's accuracy history."),
        ("Streak-Factored Habit Scoring", "The 35/25/20/20 weighted habit score now includes a streak bonus, capped so no single factor dominates."),
        ("Real-Data Web Dashboard", "The web dashboard now consumes the live habit history and recommendation APIs; the old static mock chart data was removed."),
    ],
    "22/22 M1 + 28/28 M2 smoke tests passed",
    "I reran both the Milestone 1 and Milestone 2 smoke tests after the Milestone 3 changes to prove nothing regressed, and separately verified the new Milestone 3 behaviour through the live API and the ML endpoint:",
    [
        "Milestone 1 smoke test (smoke_m1.py) - all 22 checks passed",
        "Milestone 2 smoke test (smoke_m2.py) - all 28 checks passed",
        "Adaptive difficulty model verified through ml-status (model_loaded=true)",
        "Recommendation endpoint returned sleep / wake-up / habit / challenge suggestions",
        "Habit score included the streak factor and returned correctly",
        "Web production build passed with live-data charts",
    ],
    "Other verification results:",
    [
        "The trained difficulty model ships in the repository and can be regenerated with scripts/train_difficulty_model.py",
        "Analytics are computed from real challenge attempts and alarm behaviour",
        "Recommendations were priority-sorted and returned for all four categories",
    ],
    "No Milestone 1 / Milestone 2 regression",
    "Both earlier smoke suites passed unchanged, so Milestone 3 added intelligence without breaking the authentication, alarm, or challenge workflows that came before it.",
    "User / Mobile App  >  FastAPI  >  Behavioural Analytics  >  Recent Challenge Attempts  >  Adaptive Difficulty Prediction  >  Next Challenge Selection  >  Habit Score + Streak  >  Personalised Recommendations  >  Response to User",
    [
        "The system collects the user's challenge attempts, snooze behaviour, and wake-up patterns.",
        "The analytics service converts this raw data into real sleep, wake-up, and performance values.",
        "The adaptive difficulty model (PyTorch) predicts the next challenge level from success rate, time, snoozes, and streak.",
        "The habit score is computed with the 35/25/20/20 weighting plus a streak bonus.",
        "The recommendation engine produces four priority-sorted suggestions - sleep, wake-up, habit, and challenge.",
        "The results are returned to the web dashboard and mobile app, and the next challenge is selected with the predicted difficulty.",
    ],
    "The important point: everything the user sees in Milestone 3 comes from real, verified behaviour - not static or placeholder data.",
    [
        "Placeholder analytics - the sleep and wake-up values used to be hardcoded; they are now computed from real alarm and challenge data.",
        "Missing ML inference - the adaptive difficulty model was not wired into the flow; it is now trained, loaded, and used via ml-status.",
        "Static habit score - the daily score ignored streaks; a streak factor was added and capped so it stays balanced.",
        "No recommendations - there was no coaching layer; a four-category priority-sorted recommendation engine was added.",
        "Web mock charts - the dashboard plotted static data; it now consumes the live history and recommendation APIs.",
    ],
    # --- trainer script ---
    [
        ("Opening", [
            "Good morning/afternoon. I have completed Milestone 3 of my Intelligent Cognitive Alarm Platform, and I will walk you through what I built and how I verified it.",
            "After completing Milestone 2, where the alarm required a real, server-verified cognitive challenge, I moved to Milestone 3.",
        ]),
        ("Purpose", [
            "The main focus of Milestone 3 was to make the platform smart. Now that we have verified wake-up behaviour and challenge performance, Milestone 3 turns that data into adaptive difficulty, better habit scoring, and personalised recommendations.",
        ]),
        ("What I built", [
            "I built three main things. First, an analytics and recommendation engine that computes sleep, wake-up, habit, and challenge suggestions from real data. Second, an adaptive difficulty model - a trained PyTorch model that picks the right challenge level for each user. Third, I improved habit scoring with a streak factor and wired the web dashboard to live data.",
        ]),
        ("Adaptive intelligence", [
            "The key idea is that the system learns from each user. Based on how often you solve challenges correctly, how fast you solve them, how much you snooze, and your streak, the model adjusts the difficulty. If you are improving, the challenge gets harder; if you are struggling, it gets easier. There is also a rule-based fallback so the system always works, even without the heavy model.",
        ]),
        ("Recommendations", [
            "The recommendation engine returns four categories - sleep, wake-up, habit, and challenge - sorted by priority. For example, it can suggest a 15-minute earlier bedtime, an optimal wake-up time, a habit improvement when your completion drops, and a challenge type you have not tried yet.",
        ]),
        ("Technology", [
            "For technology, I used the same stack as before - Python with FastAPI, React for the web, React Native for mobile, PostgreSQL for data, JWT with role-based access - and I added PyTorch for the ML model.",
        ]),
        ("Verification", [
            "To verify Milestone 3, I reran both the Milestone 1 and Milestone 2 smoke tests - all 22 and 28 checks still passed. I also confirmed the ML model is loaded and working through the ml-status endpoint, and that the recommendation API returns all four categories, and that the web build passes with live-data charts.",
        ]),
        ("Milestone result", [
            "So the main outcome of Milestone 3 is that the platform stopped being a fixed alarm system and became an adaptive one. It now understands each user's wake-up behaviour, adapts challenge difficulty, scores habits more meaningfully, and gives personalised coaching - all from real data.",
        ]),
        ("Next milestone", [
            "For the next milestone, I will focus on analytics, testing, and deployment - hardening the system with a proper test suite, CI/CD, structured error handling, Docker deployment, admin dashboards and reports, and mobile test coverage.",
        ]),
        ("Closing", [
            "To summarise, Milestone 3 is complete and verified. Adaptive intelligence and recommendations work end to end, and I am ready to move on to Milestone 4. Thank you.",
        ]),
    ],
)


# ---------------------------------------------------------------------------
# MILESTONE 4 : Analytics, Testing & Deployment (Weeks 7-8)
# ---------------------------------------------------------------------------
build(
    4,
    "Intelligent Cognitive Alarm Platform  |  Infosys Virtual Internship  |  Milestone 4 (Weeks 7-8)",
    [
        "Milestone 1 built the foundation, Milestone 2 made the alarm require a real server-verified cognitive challenge, and Milestone 3 added adaptive intelligence and personalised recommendations.",
        "In Milestone 4, the platform becomes production-shaped. The goal was to prove the system is robust and deployable: a proper automated test suite, a CI pipeline, structured error handling, a working Docker deployment, an admin dashboard with report exports, thorough frontend and mobile test coverage, and a final presentation. This milestone turns a working prototype into a tested, deployable product.",
    ],
    "Important libraries and frameworks actually used at this stage:",
    [
        ("Main Language", "Python", "Core backend / application logic"),
        ("Backend", "FastAPI", "API and server-side services"),
        ("Web Frontend", "React.js", "Web dashboard and admin console"),
        ("Mobile", "React Native (Expo)", "Mobile application and offline sync"),
        ("Database", "PostgreSQL", "Persistent application data (Docker)"),
        ("Testing", "pytest", "Backend unit and integration tests"),
        ("Linting", "Ruff", "Backend code quality checks"),
        ("Reports", "reportlab / openpyxl", "PDF and Excel report export"),
        ("CI/CD", "GitHub Actions", "Automated quality gates on every push"),
    ],
    [
        ("pytest", "Backend test suite (32 tests) covering auth, alarms, challenges, analytics, sync, and reports"),
        ("Ruff", "Backend linter (E/F/W rule set, clean)"),
        ("Vitest / Testing Library", "Web frontend tests (8) - components and utilities"),
        ("jest-expo", "Mobile tests (31) - sync engine, offline store, push registration, hooks"),
        ("reportlab", "PDF report generation"),
        ("openpyxl", "Excel report generation"),
        ("PyTorch", "Adaptive difficulty model (with rule-based fallback in the lean runtime)"),
        ("SQLAlchemy", "ORM for models and queries"),
        ("Alembic", "Database schema migrations"),
    ],
    "Milestone 4 - Analytics, Testing & Deployment (Weeks 7-8).",
    [
        "Build executive dashboards",
        "Add reports and visualization modules",
        "Implement testing and validations",
        "Deploy platform using Docker and cloud services",
        "Prepare final documentation and presentation",
    ],
    "Expected outcomes: a tested, deployable platform with an admin dashboard, report export, a working Docker deployment, and a final presentation. Core deliverables are complete; real cloud deployment remains a clearly-scoped follow-up.",
    "The table below lists what was actually implemented and verified for Milestone 4. The core is complete and production-shaped; the remaining items are clearly marked so nothing unimplemented is presented as done.",
    [
        ("Automated Test Suite", "Backend pytest suite grown to 32 tests (auth, alarms, challenges, analytics, sync, reports), verified on both Python 3.11 and 3.14."),
        ("Structured Error Handling", "All errors now return a consistent envelope - {detail, error_code, timestamp, status_code} - with a 422 ERR_VALIDATION code and rate-limit support."),
        ("Continuous Integration", "GitHub Actions CI runs ruff lint, the full pytest suite, both smoke tests, the web build and tests, and the mobile test suite on every push."),
        ("Docker Deployment", "Containerised deployment fixed and verified: Postgres service, a migrations entrypoint that applies alembic before startup, a lean runtime image, and an API healthcheck."),
        ("Lean Runtime", "The runtime image excludes heavy PyTorch; difficulty selection falls back to the rule-based model so the container stays small yet fully functional."),
        ("Admin Dashboard", "A new admin UI page (/admin, RBAC-gated) shows platform metrics, top challenge types, and the user list, protected so normal users get a 403."),
        ("Report Export", "PDF and Excel export added to the habit, wake-up, and challenge report endpoints, alongside the existing CSV and JSON formats, plus async report generation with a status endpoint."),
        ("Error Envelope Standard", "The structured error format is implemented additively via global handlers, so all existing clients and smoke tests keep working."),
        ("Frontend Test Coverage", "Web Vitest suite (8 tests) was added; it caught and fixed a real bug where dashboard charts were silently plotting empty data due to a field-name casing mismatch."),
        ("Mobile Test Coverage", "Mobile jest-expo suite (31 tests) covers the offline sync engine, offline store, push registration, API utilities, and hooks."),
        ("Offline Sync & Push", "The mobile app now caches data offline with a sync engine and registers its push token on login (notifications/register)."),
        ("Architecture Alignment", "Service layer modularised into domain modules, an ai_engine package added, an optional Redis adapter, and the gateway hardened with rate limiting and structured logging."),
    ],
    "32 backend + 8 web + 31 mobile tests passed",
    "Milestone 4 verification is layered - unit tests, smoke tests, a live Docker end-to-end check, and CI. In particular:",
    [
        "Backend pytest suite - 32 tests passed (verified on Python 3.11 and 3.14)",
        "Web Vitest suite - 8 tests passed (components and utilities)",
        "Mobile jest-expo suite - 31 tests passed (offline sync, store, push, hooks)",
        "Milestone 1 smoke test - all 22 checks passed",
        "Milestone 2 smoke test - all 28 checks passed",
        "Ruff lint clean across the backend",
    ],
    "Docker end-to-end verified",
    [
        "Postgres container healthy",
        "Migrations applied automatically at container startup (alembic upgrade head)",
        "Full API flow - register, login, challenge, habit, analytics, recommendations - returned 200 over HTTP",
        "PDF (%PDF) and Excel (PK) exports confirmed inside the running container",
        "Fully-local mode also works: a fresh SQLite database auto-initialises and seeds without a live server",
    ],
    "CI and mobile regression",
    "GitHub Actions runs all of the above on every push, and the Milestone 1 and Milestone 2 smoke suites still pass, so nothing was broken by the Milestone 4 changes. The final capstone presentation was produced and included in the deliverable.",
    "User / Mobile App or Web  >  FastAPI API  >  Gateway (rate limit + logging)  >  Domain Services  >  PostgreSQL  >  Structured Response / Report / Admin View",
    [
        "The user hits the API from the mobile app or web dashboard; the gateway assigns a request ID and enforces rate limiting.",
        "The API router delegates to a domain service (alarm, challenge, habit, analytics, report, admin).",
        "The service reads and writes to PostgreSQL through SQLAlchemy; migrations are managed by Alembic.",
        "Errors are returned with a consistent structured envelope.",
        "Admin users can open the dashboard, view platform metrics, and export habit / wake-up / challenge reports as CSV, JSON, PDF, or Excel, including async generation.",
        "The mobile app works offline via a local cache and sync engine, and registers its push token on login.",
        "The entire stack is containerised with Docker and runs the same quality gates in CI on every push.",
    ],
    "The important point: Milestone 4 turns the working prototype into a tested, deployable product - with automated quality gates, structured errors, Docker deployment, admin dashboards, report exports, and verified frontend and mobile coverage.",
    [
        "No error envelope - errors used to vary in shape; a global handler now returns a consistent {detail, error_code, timestamp, status_code}.",
        "Zero tests / no CI - there was no automated suite; pytest, web Vitest, and mobile jest-expo coverage plus GitHub Actions CI were added.",
        "Docker not deployable - the image and command were wrong; fixed with a root build context, a migrations entrypoint, a lean runtime, and verified compose startup.",
        "Empty dashboard charts - Vitest caught a bug where charts keyed series by camelCase instead of the snake_case API contract; fixed by mapping via dataKey.",
        "No admin view - added an RBAC-gated admin UI page with platform metrics and user management.",
        "No PDF/Excel reports - added reportlab/openpyxl export to all report formats plus async generation.",
        "Offline and push were dead ends - the mobile app never called the sync or notification APIs; wired an offline sync engine and push-token registration.",
    ],
    # --- trainer script ---
    [
        ("Opening", [
            "Good morning/afternoon. I have completed Milestone 4 of my Intelligent Cognitive Alarm Platform, and I will walk you through what I built and how I verified it.",
            "After Milestones 1, 2, and 3 established the foundation, the cognitive challenge flow, and adaptive intelligence, the final milestone focuses on analytics, testing, and deployment.",
        ]),
        ("Purpose", [
            "The main focus of Milestone 4 was to turn the working prototype into a tested, deployable product. This means a real automated test suite, a CI pipeline, structured error handling, a working Docker deployment, an admin dashboard with report exports, and thorough test coverage on the web and mobile sides.",
        ]),
        ("What I built", [
            "I added a backend test suite that grew to 32 tests, structured error handling so every error has a consistent shape, and GitHub Actions CI that runs linting, tests, and smoke tests on every push. I fixed and verified the Docker deployment - Postgres, an auto-migrations entrypoint, and a lean runtime image. I added an admin dashboard and PDF/Excel report export, plus Vitest coverage on the web and jest-expo coverage on the mobile app, and wired the mobile offline sync and push registration.",
        ]),
        ("Quality first", [
            "The key idea for Milestone 4 was that quality is a deliverable, not an afterthought. Every claim is backed by a passing test. In total we have 32 backend tests, 8 web tests, and 31 mobile tests, plus the two smoke suites from earlier milestones - all running in CI on every push. The web tests actually caught a real bug where dashboard charts were silently plotting empty data because of a field-name casing mismatch.",
        ]),
        ("Deployment", [
            "For deployment, I containerised the whole stack with Docker. The database runs as a Postgres container, migrations apply automatically at startup, and the API runs in a lean image that excludes the heavy ML dependency - difficulty selection falls back to a rule-based model. I verified the full API flow over HTTP inside the running container, and confirmed the PDF and Excel exports are produced correctly.",
        ]),
        ("Offline and push", [
            "On the mobile side, I wired the offline sync engine and push registration. The app caches alarms, challenges, and habit scores, falls back to that cache when offline, and syncs when it reconnects. It also registers its push token on login.",
        ]),
        ("Verification", [
            "To verify Milestone 4, I ran the full backend pytest suite - 32 tests - on both Python 3.11 and 3.14, the web Vitest suite, the mobile jest suite, and both earlier smoke suites. Everything passed, and the Docker deployment was verified end to end.",
        ]),
        ("Milestone result", [
            "So the main outcome of Milestone 4 is that the intelligent cognitive alarm platform is now production-shaped: tested, linted, containerised, with an admin dashboard, report exports, and CI running on every push. The core is complete.",
        ]),
        ("Next steps", [
            "The remaining, clearly-scoped follow-ups are the real cloud deployment, a native push-delivery provider (FCM/APNs) with on-device local alarm scheduling, and broader end-to-end flow coverage. On-device push needs a development build, since Expo Go disables native push. Nothing unimplemented is presented as done.",
        ]),
        ("Closing", [
            "To summarise, Milestone 4 is complete for its core scope and verified. The platform is tested, deployable, and ready for the next steps. Thank you.",
        ]),
    ],
)
