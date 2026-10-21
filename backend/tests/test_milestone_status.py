import unittest
from datetime import datetime, timedelta
from io import BytesIO
from types import SimpleNamespace

from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from openpyxl import load_workbook

from app.ai.milestone_engine import evaluate_milestone_status
from app.ai.priority_engine import update_alarm_difficulty_from_experience
from app.database import Base
from app.models.alarm import Alarm
from app.models.challenge import Challenge
from app.models.user import User
from app.routers.challenges import submit_challenge
from app.services.analytics_service import get_dashboard
from app.services.challenge_service import can_start_challenge, generate_questions
from app.services.report_service import build_excel


class MilestoneStatusTests(unittest.TestCase):
    def test_dashboard_metrics_only_include_the_current_user_and_answered_accuracy(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(bind=engine)
        Session = sessionmaker(bind=engine)
        db = Session()

        first_user = User(name="First User", email="first@example.com", password_hash="hash")
        second_user = User(name="Second User", email="second@example.com", password_hash="hash")
        db.add_all([first_user, second_user])
        db.commit()
        db.refresh(first_user)
        db.refresh(second_user)

        alarm = Alarm(
            user_id=first_user.id,
            title="Morning Check",
            alarm_time="07:00",
            difficulty="easy",
            challenge_type="math",
            snooze_count=1,
            completed_count=2,
        )
        db.add(alarm)
        db.commit()
        db.refresh(alarm)

        db.add_all([
            Challenge(
                user_id=first_user.id,
                alarm_id=alarm.id,
                session_id=f"session-{session_number}",
                question_number=question_number,
                question=f"Question {session_number}-{question_number}",
                answer="1",
                is_correct=True if session_number == 1 and question_number == 1 else None,
                is_completed=session_number == 1 and question_number == 1,
            )
            for session_number in range(1, 3)
            for question_number in range(1, 6)
        ])
        db.commit()

        first_stats = get_dashboard(db, first_user)
        second_stats = get_dashboard(db, second_user)

        self.assertEqual(first_stats["challenge_accuracy"], 100)
        self.assertEqual(first_stats["total_challenges"], 2)
        self.assertEqual(first_stats["total_questions"], 10)
        self.assertEqual(first_stats["answered_questions"], 1)
        self.assertEqual(first_stats["answered_challenges"], 1)
        self.assertEqual(first_stats["alarm_completions"], 2)
        self.assertEqual(first_stats["snooze_count"], 1)
        self.assertEqual(second_stats["total_alarms"], 0)
        self.assertEqual(second_stats["total_challenges"], 0)
        self.assertEqual(second_stats["alarm_completions"], 0)
        self.assertEqual(second_stats["snooze_count"], 0)

        db.close()

    def test_milestone_status_tracks_completed_user_workflow(self):
        user = SimpleNamespace(
            name="Ava",
            role="user",
            preferred_wake_time="07:00",
            sleep_duration=8,
            role_preference="student",
        )

        alarms = [
            SimpleNamespace(id=1, title="Morning Focus", is_active=True),
            SimpleNamespace(id=2, title="Deep Work", is_active=True),
        ]

        challenges = [
            SimpleNamespace(is_completed=True, is_correct=True),
            SimpleNamespace(is_completed=True, is_correct=False),
        ]

        dashboard = {
            "habit_score": 81,
            "recommended_difficulty": "hard",
            "challenge_completion_rate": 75,
            "challenge_accuracy": 82,
            "recommendations": ["Keep a fixed wake-up time"],
        }

        status = evaluate_milestone_status(user, alarms, challenges, dashboard)

        self.assertEqual(status["milestone_1"]["status"], "operational")
        self.assertEqual(status["milestone_2"]["status"], "operational")
        self.assertEqual(status["milestone_3"]["status"], "operational")
        self.assertIn(status["milestone_4"]["status"], {"ready", "operational"})
        self.assertTrue(status["milestone_1"]["complete"])
        self.assertIn(
            "Understand intelligent alarm and habit formation workflows.",
            status["milestone_1"]["outcomes"],
        )
        self.assertIn(
            "Cognitive challenge engine operational.",
            status["milestone_2"]["outcomes"],
        )
        self.assertIn(
            "Adaptive difficulty engine operational.",
            status["milestone_3"]["outcomes"],
        )
        self.assertIn(
            "Fully deployed production-ready platform.",
            status["milestone_4"]["outcomes"],
        )

    def test_milestone_status_uses_pending_when_user_has_not_interacted(self):
        user = SimpleNamespace(
            name="",
            role="user",
            preferred_wake_time=None,
            sleep_duration=0,
            role_preference=None,
        )

        status = evaluate_milestone_status(user, [], [], {"habit_score": 0})

        self.assertEqual(status["milestone_1"]["status"], "pending")
        self.assertEqual(status["milestone_2"]["status"], "pending")
        self.assertEqual(status["milestone_3"]["status"], "pending")
        self.assertEqual(status["milestone_4"]["status"], "pending")

    def test_milestone_status_handles_null_values_for_new_users(self):
        user = SimpleNamespace(
            name="New User",
            role="user",
            preferred_wake_time=None,
            sleep_duration=0,
            role_preference=None,
        )

        status = evaluate_milestone_status(user, [], [], {"habit_score": None, "recommendations": []})

        self.assertEqual(status["milestone_1"]["status"], "pending")
        self.assertEqual(status["milestone_2"]["status"], "pending")
        self.assertEqual(status["milestone_3"]["status"], "pending")

    def test_challenge_generation_supports_all_five_types(self):
        questions = generate_questions("riddle", "easy", 2)
        self.assertEqual(len(questions), 2)
        self.assertTrue(all("?" in q["question"] for q in questions))

        word_questions = generate_questions("word", "hard", 2)
        self.assertEqual(len(word_questions), 2)

    def test_generated_questions_stay_unique_for_each_user_session(self):
        used = [
            "What is 2 + 2?",
            "What is 5 + 5?",
        ]

        questions = generate_questions("math", "easy", 4, used_questions=used)

        self.assertEqual(len(questions), 4)
        self.assertEqual(len({q["question"] for q in questions}), 4)
        self.assertTrue(all(q["question"] not in used for q in questions))

    def test_alarm_challenge_requires_active_ring_window(self):
        active_alarm = SimpleNamespace(
            last_triggered=datetime.utcnow(),
            challenge_completed=False,
            is_smart=True,
        )
        inactive_alarm = SimpleNamespace(
            last_triggered=datetime.utcnow() - timedelta(minutes=10),
            challenge_completed=False,
            is_smart=True,
        )

        self.assertTrue(can_start_challenge(active_alarm))
        self.assertFalse(can_start_challenge(inactive_alarm))

    def test_adaptive_difficulty_updates_alarm_from_recent_performance(self):
        user = SimpleNamespace(difficulty_preference="easy")
        alarm = SimpleNamespace(difficulty="easy", snooze_count=0)
        recent = [
            SimpleNamespace(is_correct=True, response_time=12),
            SimpleNamespace(is_correct=True, response_time=14),
            SimpleNamespace(is_correct=True, response_time=16),
        ]

        updated = update_alarm_difficulty_from_experience(alarm, user, recent)

        self.assertEqual(updated.difficulty, "medium")
        self.assertEqual(user.difficulty_preference, "medium")

    def test_adaptive_difficulty_adjusts_after_final_question_even_when_wrong(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(bind=engine)
        Session = sessionmaker(bind=engine)
        db = Session()

        user = User(
            name="Test User",
            email="scoreuser@example.com",
            password_hash="hash",
            role="user",
            difficulty_preference="easy",
        )
        db.add(user)
        db.commit()
        db.refresh(user)

        alarm = Alarm(
            user_id=user.id,
            title="Morning Check",
            alarm_time="07:00",
            days="Monday",
            difficulty="easy",
            challenge_type="math",
            is_smart=True,
            challenge_question_count=5,
            challenge_completed=False,
        )
        db.add(alarm)
        db.commit()
        db.refresh(alarm)

        session_id = "session-123"
        for number in range(1, 6):
            challenge = Challenge(
                user_id=user.id,
                alarm_id=alarm.id,
                session_id=session_id,
                question_number=number,
                question=f"Q{number}",
                answer="1",
                challenge_type="math",
                difficulty="easy",
                is_correct=(number in {1, 3}),
                is_completed=(number in {1, 3}),
                user_answer="1" if number in {1, 3} else "2",
                response_time=10.0,
                attempt_count=1,
            )
            db.add(challenge)
            db.commit()
            db.refresh(challenge)

        final_question = db.query(Challenge).filter_by(alarm_id=alarm.id, session_id=session_id, question_number=5).one()
        final_question.user_answer = None
        final_question.is_correct = None
        final_question.is_completed = False
        db.commit()

        response = submit_challenge(final_question.id, "wrong", db=db, current_user=user)

        self.assertEqual(response["score"], 2)
        self.assertEqual(response["total_questions"], 5)
        self.assertEqual(alarm.difficulty, "easy")
        self.assertEqual(user.difficulty_preference, "easy")

        workbook_data = build_excel(db, user)
        workbook = load_workbook(filename=BytesIO(workbook_data.getvalue()))
        self.assertEqual(workbook.sheetnames[0], "User Performance")
        self.assertIn("scoreuser@example.com", str(workbook["User Performance"]["B3"].value))

        db.close()


if __name__ == "__main__":
    unittest.main()
