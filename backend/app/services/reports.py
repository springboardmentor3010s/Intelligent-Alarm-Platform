"""Reports & Export: habit, wake-up, challenge performance, productivity and sleep reports as PDF or Excel."""
import io
from datetime import timedelta

import pandas as pd
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from app.models import ChallengeInstance, SleepLog, User
from app.services import analytics as A

KINDS = ("habit", "wake", "challenge", "productivity", "sleep")


def build(db: Session, user: User, kind: str, days: int):
    """Returns (title, summary_pairs, [(section_title, DataFrame)])."""
    uid = user.id
    if kind == "habit":
        s = A.compute_habit_score(db, uid)
        summary = [("Habit score", s["total"]), ("Wake consistency", s["wake_consistency"]),
                   ("Challenge completion", s["challenge_completion"]), ("Snooze reduction", s["snooze_reduction"]),
                   ("Sleep adherence", s["sleep_adherence"])]
        return "Habit report", summary, [("Habit score history", pd.DataFrame(A.habit_history(db, uid, days)))]
    if kind == "wake":
        w = A.wake_stats(db, uid, days)
        rows = [{"date": x.local_date.isoformat(), "alarm": x.label, "status": x.status, "snoozes": x.snooze_count,
                 "wrong_answers": x.wrong_count, "method": x.verification_method, "alertness": x.alertness_score,
                 "seconds_to_dismiss": round((x.dismissed_at - x.started_at).total_seconds()) if x.dismissed_at else None}
                for x in A.finished_sessions(db, uid, days)]
        summary = [("Alarms", w["total"]), ("Dismissed", w["dismissed"]), ("Missed", w["missed"]),
                   ("Success rate", w["success_rate"]), ("Current streak", w["current_streak"])]
        return "Wake-up report", summary, [("Alarm history", pd.DataFrame(rows)),
                                           ("By weekday", pd.DataFrame(w["by_weekday"]))]
    if kind == "challenge":
        c = A.challenge_stats(db, uid, days)
        since = A.today_of(db, uid) - timedelta(days=days)
        inst = (db.query(ChallengeInstance).filter(ChallengeInstance.user_id == uid, ChallengeInstance.is_practice.is_(False),
                                                   ChallengeInstance.answered_at.isnot(None),
                                                   ChallengeInstance.issued_at >= since).all())
        log = [{"when": i.issued_at.strftime("%Y-%m-%d %H:%M"), "type": i.challenge_type, "difficulty": i.difficulty,
                "correct": i.is_correct, "seconds": i.time_taken_seconds, "timed_out": i.timed_out} for i in inst]
        return "Challenge performance report", [("Attempts", c["total"]), ("Accuracy", c["accuracy"]),
                                                ("Avg seconds", c["avg_seconds"])], \
            [("By type", pd.DataFrame(c["by_type"])), ("By difficulty", pd.DataFrame(c["by_difficulty"])),
             ("Attempt log", pd.DataFrame(log))]
    if kind == "productivity":
        p, sl = A.productivity_correlation(db, uid, days), A.sleep_stats(db, uid, days)
        return "Productivity report", [("Days analysed", p["n"]), ("Avg productivity (1-5)", p.get("avg_productivity"))], \
            [("Correlations", pd.DataFrame(p["correlations"])), ("Daily ratings", pd.DataFrame(sl["series"]))]
    if kind == "sleep":
        s = A.sleep_stats(db, uid, days)
        summary = [("Nights logged", s["n"]), ("Avg duration (min)", s.get("avg_duration_minutes")),
                   ("Avg bedtime", s.get("avg_bedtime")), ("Sleep debt (min)", s.get("sleep_debt_minutes"))]
        return "Sleep analytics report", summary, [("Sleep log", pd.DataFrame(s["series"]))]
    raise ValueError("unknown report kind")


def _clean(df: pd.DataFrame) -> pd.DataFrame:
    return df if len(df.columns) else pd.DataFrame({"info": ["No data for this period"]})


def to_excel(title, summary, sections) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        pd.DataFrame(summary, columns=["Metric", "Value"]).to_excel(xw, sheet_name="Summary", index=False)
        for name, df in sections:
            _clean(df).to_excel(xw, sheet_name=name[:31], index=False)
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", start_color="1B2142")
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(40, max(len(str(c.value or "")) for c in col) + 3)
    return buf.getvalue()


def _table(rows, header_bg="#1B2142"):
    t = Table(rows, hAlign="LEFT", repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(header_bg)), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCCCCC")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F5FA")])]))
    return t


def to_pdf(title, summary, sections, user: User, days: int) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.8 * cm, bottomMargin=1.8 * cm, title=title)
    st = getSampleStyleSheet()
    els = [Paragraph(f"Wakeful — {title}", st["Title"]),
           Paragraph(f"{user.full_name} · {user.email} · last {days} days", st["Normal"]), Spacer(1, 14),
           Paragraph("Summary", st["Heading2"]),
           _table([["Metric", "Value"]] + [[k, "-" if v is None else str(v)] for k, v in summary])]
    for name, df in sections:
        els += [Spacer(1, 14), Paragraph(name, st["Heading2"])]
        df = _clean(df).head(45)
        els.append(_table([list(map(str, df.columns))] + [["-" if pd.isna(v) else str(v) for v in r] for r in df.values.tolist()]))
    doc.build(els)
    return buf.getvalue()
