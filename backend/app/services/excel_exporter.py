import io
from typing import Any, Dict
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def apply_header_style(cell):
    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="0D172C", end_color="0D172C", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def apply_sub_header_style(cell):
    cell.font = Font(name="Calibri", size=11, bold=True, color="1E293B")
    cell.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    cell.alignment = Alignment(horizontal="left", vertical="center")


def apply_cell_style(cell, bold=False, align="left"):
    cell.font = Font(name="Calibri", size=10, bold=bold, color="0F172A")
    cell.alignment = Alignment(horizontal=align, vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )
    cell.border = thin_border


def auto_fit_columns(ws, min_width=15, max_width=45):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = cell.value
            if val is not None:
                lines = str(val).split("\n")
                line_max = max(len(l) for l in lines)
                if line_max > max_len:
                    max_len = line_max
        ws.column_dimensions[col_letter].width = max(min(max_len + 3, max_width), min_width)


def generate_excel_report(data: Dict[str, Any]) -> io.BytesIO:
    """
    Generates a structured multi-sheet Excel workbook from the aggregated report data.
    Returns an in-memory BytesIO buffer containing the .xlsx bytes.
    """
    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    is_demo = data.get("is_demo", False)

    # -------------------------------------------------------------
    # Sheet 1: Summary
    # -------------------------------------------------------------
    ws_summary = wb.create_sheet(title="Summary")
    ws_summary.views.sheetView[0].showGridLines = True

    row = 1
    if is_demo:
        ws_summary.merge_cells("A1:D1")
        demo_cell = ws_summary["A1"]
        demo_cell.value = "DEMO REPORT — SAMPLE DATA (READ-ONLY)"
        demo_cell.font = Font(name="Calibri", size=11, bold=True, color="92400E")
        demo_cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        demo_cell.alignment = Alignment(horizontal="center", vertical="center")
        row = 3

    # Title
    ws_summary.cell(row=row, column=1, value="INTELLIGENT COGNITIVE ALARM — PERFORMANCE REPORT").font = Font(
        name="Calibri", size=14, bold=True, color="0F172A"
    )
    row += 2

    # Metadata
    metadata_rows = [
        ("User Name", data.get("user_name", "User")),
        ("User Email", data.get("user_email", "N/A")),
        ("Generated At", data.get("generated_at", "N/A")),
        ("Report Type", "Milestone 4 — Phase 1 Performance Report"),
    ]
    for label, val in metadata_rows:
        c1 = ws_summary.cell(row=row, column=1, value=label)
        c2 = ws_summary.cell(row=row, column=2, value=val)
        apply_cell_style(c1, bold=True)
        apply_cell_style(c2)
        row += 1

    row += 1
    # Key Highlights Table
    ws_summary.cell(row=row, column=1, value="Metric Category")
    ws_summary.cell(row=row, column=2, value="Key Indicator")
    ws_summary.cell(row=row, column=3, value="Current Value")
    for col in range(1, 4):
        apply_header_style(ws_summary.cell(row=row, column=col))
    row += 1

    habit = data.get("habit_report", {})
    wake = data.get("wake_up_report", {})
    challenge = data.get("challenge_report", {})
    sleep = data.get("sleep_adherence_report", {})
    prod = data.get("productivity_report", {})

    kpi_rows = [
        ("Habit", "Overall Habit Score", f"{habit.get('overall_habit_score', 0):.1f} / 100"),
        ("Habit", "Wake-Up Consistency", f"{habit.get('wake_up_consistency', 0):.1f}%"),
        ("Wake-Up", "Total Sessions", str(wake.get("total_sessions", 0))),
        ("Wake-Up", "Successful Dismissals", str(wake.get("successful_dismissals", 0))),
        ("Cognitive Challenges", "Completion Percentage", f"{challenge.get('completion_percentage', 0):.1f}%"),
        ("Cognitive Challenges", "Accuracy Percentage", f"{challenge.get('accuracy_percentage', 0):.1f}%"),
        ("Sleep Adherence", "Preferred Wake Time", str(sleep.get("preferred_wake_time", "Not available"))),
        ("Sleep Adherence", "Schedule Adherence", f"{sleep.get('sleep_schedule_adherence', 0):.1f}%"),
        ("Productivity", "Status", str(prod.get("status", "Not available"))),
    ]

    for cat, kpi, val in kpi_rows:
        c1 = ws_summary.cell(row=row, column=1, value=cat)
        c2 = ws_summary.cell(row=row, column=2, value=kpi)
        c3 = ws_summary.cell(row=row, column=3, value=val)
        apply_cell_style(c1)
        apply_cell_style(c2, bold=True)
        apply_cell_style(c3, align="right")
        row += 1

    auto_fit_columns(ws_summary)

    # -------------------------------------------------------------
    # Sheet 2: Habit Metrics
    # -------------------------------------------------------------
    ws_habit = wb.create_sheet(title="Habit Metrics")
    ws_habit.views.sheetView[0].showGridLines = True

    h_headers = ["Habit Metric", "Value", "Unit / Scale"]
    for idx, h in enumerate(h_headers, start=1):
        c = ws_habit.cell(row=1, column=idx, value=h)
        apply_header_style(c)

    habit_items = [
        ("Overall Habit Score", habit.get("overall_habit_score", 0), "Score (0 - 100)"),
        ("Wake-Up Consistency", habit.get("wake_up_consistency", 0), "Percentage (%)"),
        ("Challenge Completion", habit.get("challenge_completion", 0), "Percentage (%)"),
        ("Snooze Reduction", habit.get("snooze_reduction", 0), "Score (0 - 100)"),
        ("Sleep Schedule Adherence", habit.get("sleep_schedule_adherence", 0), "Percentage (%)"),
        ("Sessions Evaluated", habit.get("sessions_evaluated", 0), "Count"),
        ("Average Snooze Count", habit.get("average_snooze_count", 0), "Snoozes per session"),
    ]

    for r_idx, (metric, val, unit) in enumerate(habit_items, start=2):
        c1 = ws_habit.cell(row=r_idx, column=1, value=metric)
        c2 = ws_habit.cell(row=r_idx, column=2, value=val)
        c3 = ws_habit.cell(row=r_idx, column=3, value=unit)
        apply_cell_style(c1, bold=True)
        apply_cell_style(c2, align="right")
        apply_cell_style(c3)

    auto_fit_columns(ws_habit)

    # -------------------------------------------------------------
    # Sheet 3: Wake-Up Performance
    # -------------------------------------------------------------
    ws_wake = wb.create_sheet(title="Wake-Up Performance")
    ws_wake.views.sheetView[0].showGridLines = True

    w_headers = ["Wake-Up Performance Metric", "Value"]
    for idx, h in enumerate(w_headers, start=1):
        c = ws_wake.cell(row=1, column=idx, value=h)
        apply_header_style(c)

    wake_items = [
        ("Total Sessions", wake.get("total_sessions", 0)),
        ("Wake-Up Consistency (%)", f"{wake.get('wake_up_consistency', 0):.1f}%"),
        ("Average Response Time (seconds)", wake.get("average_response_time", 0)),
        ("Successful Dismissals", wake.get("successful_dismissals", 0)),
        ("Total Snoozes Used", wake.get("total_snoozes", 0)),
    ]

    for r_idx, (metric, val) in enumerate(wake_items, start=2):
        c1 = ws_wake.cell(row=r_idx, column=1, value=metric)
        c2 = ws_wake.cell(row=r_idx, column=2, value=val)
        apply_cell_style(c1, bold=True)
        apply_cell_style(c2, align="right")

    r_next = len(wake_items) + 3
    # Status Breakdown table
    c_sub = ws_wake.cell(row=r_next, column=1, value="Session Status Breakdown")
    c_sub2 = ws_wake.cell(row=r_next, column=2, value="Count")
    apply_sub_header_style(c_sub)
    apply_sub_header_style(c_sub2)
    r_next += 1

    status_bd = wake.get("status_breakdown", {})
    if status_bd:
        for status_name, cnt in sorted(status_bd.items()):
            c1 = ws_wake.cell(row=r_next, column=1, value=status_name.capitalize())
            c2 = ws_wake.cell(row=r_next, column=2, value=cnt)
            apply_cell_style(c1)
            apply_cell_style(c2, align="right")
            r_next += 1
    else:
        c1 = ws_wake.cell(row=r_next, column=1, value="No session history")
        c2 = ws_wake.cell(row=r_next, column=2, value=0)
        apply_cell_style(c1)
        apply_cell_style(c2, align="right")

    auto_fit_columns(ws_wake)

    # -------------------------------------------------------------
    # Sheet 4: Cognitive Challenges
    # -------------------------------------------------------------
    ws_chal = wb.create_sheet(title="Cognitive Challenges")
    ws_chal.views.sheetView[0].showGridLines = True

    c_headers = ["Challenge Overview Metric", "Value"]
    for idx, h in enumerate(c_headers, start=1):
        c = ws_chal.cell(row=1, column=idx, value=h)
        apply_header_style(c)

    chal_items = [
        ("Total Challenges", challenge.get("total_challenges", 0)),
        ("Completed Challenges", challenge.get("completed_challenges", 0)),
        ("Correct Answers", challenge.get("correct_answers", 0)),
        ("Incorrect Answers", challenge.get("incorrect_answers", 0)),
        ("Completion Percentage", f"{challenge.get('completion_percentage', 0):.1f}%"),
        ("Accuracy Percentage", f"{challenge.get('accuracy_percentage', 0):.1f}%"),
        ("Average Response Time (seconds)", challenge.get("average_response_time", 0)),
    ]

    for r_idx, (metric, val) in enumerate(chal_items, start=2):
        c1 = ws_chal.cell(row=r_idx, column=1, value=metric)
        c2 = ws_chal.cell(row=r_idx, column=2, value=val)
        apply_cell_style(c1, bold=True)
        apply_cell_style(c2, align="right")

    # Table: Type Breakdown
    r_idx = len(chal_items) + 3
    type_headers = ["Challenge Type", "Total", "Correct", "Incorrect", "Accuracy (%)"]
    for c_i, th in enumerate(type_headers, start=1):
        c = ws_chal.cell(row=r_idx, column=c_i, value=th)
        apply_header_style(c)
    r_idx += 1

    types = challenge.get("challenge_types", [])
    if types:
        for t in types:
            c1 = ws_chal.cell(row=r_idx, column=1, value=t.get("challenge_type", ""))
            c2 = ws_chal.cell(row=r_idx, column=2, value=t.get("total", 0))
            c3 = ws_chal.cell(row=r_idx, column=3, value=t.get("correct", 0))
            c4 = ws_chal.cell(row=r_idx, column=4, value=t.get("incorrect", 0))
            c5 = ws_chal.cell(row=r_idx, column=5, value=f"{t.get('accuracy', 0):.1f}%")
            apply_cell_style(c1, bold=True)
            apply_cell_style(c2, align="right")
            apply_cell_style(c3, align="right")
            apply_cell_style(c4, align="right")
            apply_cell_style(c5, align="right")
            r_idx += 1
    else:
        for c_i in range(1, 6):
            c = ws_chal.cell(row=r_idx, column=c_i, value="No data" if c_i == 1 else "-")
            apply_cell_style(c)
        r_idx += 1

    # Table: Difficulty Breakdown
    r_idx += 1
    diff_headers = ["Difficulty Level", "Total", "Correct", "Incorrect", "Accuracy (%)"]
    for c_i, dh in enumerate(diff_headers, start=1):
        c = ws_chal.cell(row=r_idx, column=c_i, value=dh)
        apply_header_style(c)
    r_idx += 1

    diffs = challenge.get("difficulty_levels", [])
    if diffs:
        for d in diffs:
            c1 = ws_chal.cell(row=r_idx, column=1, value=d.get("difficulty", ""))
            c2 = ws_chal.cell(row=r_idx, column=2, value=d.get("total", 0))
            c3 = ws_chal.cell(row=r_idx, column=3, value=d.get("correct", 0))
            c4 = ws_chal.cell(row=r_idx, column=4, value=d.get("incorrect", 0))
            c5 = ws_chal.cell(row=r_idx, column=5, value=f"{d.get('accuracy', 0):.1f}%")
            apply_cell_style(c1, bold=True)
            apply_cell_style(c2, align="right")
            apply_cell_style(c3, align="right")
            apply_cell_style(c4, align="right")
            apply_cell_style(c5, align="right")
            r_idx += 1
    else:
        for c_i in range(1, 6):
            c = ws_chal.cell(row=r_idx, column=c_i, value="No data" if c_i == 1 else "-")
            apply_cell_style(c)
        r_idx += 1

    auto_fit_columns(ws_chal)

    # -------------------------------------------------------------
    # Sheet 5: Sleep Adherence
    # -------------------------------------------------------------
    ws_sleep = wb.create_sheet(title="Sleep Adherence")
    ws_sleep.views.sheetView[0].showGridLines = True

    s_headers = ["Sleep & Schedule Metric", "Value"]
    for idx, h in enumerate(s_headers, start=1):
        c = ws_sleep.cell(row=1, column=idx, value=h)
        apply_header_style(c)

    sleep_items = [
        ("Preferred Wake Time", str(sleep.get("preferred_wake_time", "Not available"))),
        ("Target Sleep Duration", str(sleep.get("sleep_duration", "Not available"))),
        ("Sleep Schedule Adherence", f"{sleep.get('sleep_schedule_adherence', 0):.1f}%"),
        ("Schedule Alignment", str(sleep.get("schedule_alignment", "Not available"))),
        ("Profile Data Configured", "Yes" if sleep.get("has_profile_data") else "No"),
    ]

    for r_idx, (metric, val) in enumerate(sleep_items, start=2):
        c1 = ws_sleep.cell(row=r_idx, column=1, value=metric)
        c2 = ws_sleep.cell(row=r_idx, column=2, value=val)
        apply_cell_style(c1, bold=True)
        apply_cell_style(c2)

    auto_fit_columns(ws_sleep)

    # -------------------------------------------------------------
    # Sheet 6: Productivity
    # -------------------------------------------------------------
    ws_prod = wb.create_sheet(title="Productivity")
    ws_prod.views.sheetView[0].showGridLines = True

    p_headers = ["Productivity Field", "Details"]
    for idx, h in enumerate(p_headers, start=1):
        c = ws_prod.cell(row=1, column=idx, value=h)
        apply_header_style(c)

    prod_items = [
        ("Productivity Data Available", "Yes" if prod.get("has_productivity_data") else "No"),
        ("Status", str(prod.get("status", "Not available"))),
        ("Productivity Goals", str(prod.get("productivity_goals") or "Productivity data is not available yet.")),
        ("Message", str(prod.get("message", "Productivity data is not available yet."))),
    ]

    for r_idx, (field, val) in enumerate(prod_items, start=2):
        c1 = ws_prod.cell(row=r_idx, column=1, value=field)
        c2 = ws_prod.cell(row=r_idx, column=2, value=val)
        apply_cell_style(c1, bold=True)
        apply_cell_style(c2)

    auto_fit_columns(ws_prod)

    # Save to BytesIO
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
