"""
ReportGenerator — Module 12 (Reports & Export System)

Builds downloadable PDF and Excel habit/performance reports for a user,
combining the habit-score formula, behavioral analytics, and raw
challenge history for a given time window.
"""
from datetime import datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from app.models.models import CognitiveChallenge, User
from app.services.behavioral_analytics import calculate_behavioral_summary
from app.services.habit_scoring import calculate_habit_score


def _challenge_rows(db: Session, user_id: int, days: int) -> list[CognitiveChallenge]:
    cutoff = datetime.now() - timedelta(days=days)
    return (
        db.query(CognitiveChallenge)
        .filter(CognitiveChallenge.user_id == user_id, CognitiveChallenge.started_at >= cutoff)
        .order_by(CognitiveChallenge.started_at.desc())
        .all()
    )


def build_habit_report_pdf(db: Session, user: User, days: int = 7) -> BytesIO:
    habit = calculate_habit_score(db, user.id, days)
    behavioral = calculate_behavioral_summary(db, user.id, days)
    challenges = _challenge_rows(db, user.id, days)

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, title="Cognitive Alarm Habit Report")
    styles = getSampleStyleSheet()
    elements = [
        Paragraph("Cognitive Alarm Platform — Habit Report", styles["Title"]),
        Paragraph(f"User: {user.name} ({user.email})", styles["Normal"]),
        Paragraph(f"Period: last {days} day(s) — generated {datetime.now():%Y-%m-%d %H:%M}", styles["Normal"]),
        Spacer(1, 16),
        Paragraph(f"Overall Habit Score: {habit['habit_score']} / 100", styles["Heading2"]),
        Spacer(1, 8),
    ]

    comp_data = [["Component", "Weight", "Score", "Contribution"]]
    for comp in habit["components"].values():
        comp_data.append([
            comp["label"], f"{int(comp['weight'] * 100)}%", comp["score"], comp["contribution"],
        ])
    comp_table = Table(comp_data, hAlign="LEFT")
    comp_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4f46e5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]))
    elements += [comp_table, Spacer(1, 20)]

    elements.append(Paragraph("Behavioral Summary", styles["Heading2"]))
    beh_data = [["Metric", "Score"]]
    for key in ("wakeup_consistency", "challenge_completion", "snooze_reduction", "sleep_adherence"):
        beh_data.append([key.replace("_", " ").title(), behavioral[key]["score"]])
    beh_table = Table(beh_data, hAlign="LEFT")
    beh_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4f46e5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]))
    elements += [beh_table, Spacer(1, 20)]

    elements.append(Paragraph(f"Challenge History ({len(challenges)} record(s))", styles["Heading2"]))
    ch_data = [["Date", "Type", "Difficulty", "Correct", "Score", "Attempts"]]
    for c in challenges[:50]:
        ch_data.append([
            c.started_at.strftime("%Y-%m-%d %H:%M"),
            c.challenge_type,
            c.difficulty,
            "Yes" if c.is_correct else ("No" if c.is_correct is False else "—"),
            c.score,
            c.attempts,
        ])
    ch_table = Table(ch_data, hAlign="LEFT")
    ch_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4f46e5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
    ]))
    elements.append(ch_table)

    doc.build(elements)
    buffer.seek(0)
    return buffer


def build_habit_report_excel(db: Session, user: User, days: int = 7) -> BytesIO:
    habit = calculate_habit_score(db, user.id, days)
    behavioral = calculate_behavioral_summary(db, user.id, days)
    challenges = _challenge_rows(db, user.id, days)

    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    summary.append(["Cognitive Alarm Platform — Habit Report"])
    summary["A1"].font = Font(size=14, bold=True)
    summary.append([f"User: {user.name} ({user.email})"])
    summary.append([f"Period: last {days} day(s)"])
    summary.append([f"Overall Habit Score: {habit['habit_score']} / 100"])
    summary.append([])
    summary.append(["Component", "Weight", "Score", "Contribution"])
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)
    for comp in habit["components"].values():
        summary.append([comp["label"], f"{int(comp['weight'] * 100)}%", comp["score"], comp["contribution"]])

    summary.append([])
    summary.append(["Behavioral Metric", "Score"])
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)
    for key in ("wakeup_consistency", "challenge_completion", "snooze_reduction", "sleep_adherence"):
        summary.append([key.replace("_", " ").title(), behavioral[key]["score"]])

    for col, width in zip("ABCD", (40, 12, 10, 14)):
        summary.column_dimensions[col].width = width

    sheet = wb.create_sheet("Challenges")
    headers = ["Date", "Type", "Difficulty", "Correct", "Score", "Attempts", "Duration (s)"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for c in challenges:
        sheet.append([
            c.started_at.strftime("%Y-%m-%d %H:%M"),
            c.challenge_type,
            c.difficulty,
            "Yes" if c.is_correct else ("No" if c.is_correct is False else "—"),
            c.score,
            c.attempts,
            c.completion_duration_seconds,
        ])
    for col, width in zip("ABCDEFG", (18, 12, 12, 10, 8, 10, 14)):
        sheet.column_dimensions[col].width = width

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
